"""
특허명세서 생성기
==================

발명자료를 기반으로 완전한 특허명세서를 생성하는 핵심 클래스입니다.
"""

from datetime import datetime
from pathlib import Path
from typing import Literal

from patent_generator.models.invention import InventionData
from patent_generator.generators.sections import (
    generate_technical_field_section,
    generate_background_art_section,
    generate_problem_section,
    generate_solution_section,
    generate_effect_section,
    generate_drawings_section,
    generate_embodiments_section,
    generate_reference_signs_section,
    generate_claims_section,
    generate_abstract_section,
)


class PatentSpecificationGenerator:
    """
    특허명세서 자동 생성기

    발명자료(InventionData)를 입력받아 한국 특허명세서 형식의
    문서를 생성합니다.
    """

    def __init__(self, invention_data: InventionData):
        """
        Args:
            invention_data: 발명자료 데이터 객체
        """
        self.invention_data = invention_data
        self._specification_text: str | None = None

    @classmethod
    def from_yaml(cls, yaml_path: str) -> "PatentSpecificationGenerator":
        """YAML 파일에서 생성기 초기화"""
        invention_data = InventionData.from_yaml(yaml_path)
        return cls(invention_data)

    @classmethod
    def from_json(cls, json_path: str) -> "PatentSpecificationGenerator":
        """JSON 파일에서 생성기 초기화"""
        invention_data = InventionData.from_json(json_path)
        return cls(invention_data)

    def generate(self) -> str:
        """
        특허명세서 전체 텍스트 생성

        Returns:
            생성된 특허명세서 텍스트
        """
        sections = []

        # 헤더
        sections.append(self._generate_header())

        # 명세서 본문
        sections.append("【명세서】")
        sections.append("")

        # 발명의 명칭
        sections.append("【발명의 명칭】")
        sections.append(self.invention_data.title)
        sections.append("")

        # 기술분야
        sections.append(generate_technical_field_section(self.invention_data.technical_field))

        # 발명의 배경이 되는 기술
        sections.append(generate_background_art_section(self.invention_data.background_art))

        # 발명의 내용
        sections.append("【발명의 내용】")
        sections.append("")

        # 해결하려는 과제
        sections.append(generate_problem_section(self.invention_data.problem_to_solve))

        # 과제의 해결 수단
        sections.append(generate_solution_section(self.invention_data.solution))

        # 발명의 효과
        sections.append(generate_effect_section(self.invention_data.effect))

        # 도면의 간단한 설명
        drawings_section = generate_drawings_section(self.invention_data.drawings)
        if drawings_section:
            sections.append(drawings_section)

        # 발명을 실시하기 위한 구체적인 내용
        embodiments_section = generate_embodiments_section(self.invention_data.embodiments)
        if embodiments_section:
            sections.append(embodiments_section)

        # 부호의 설명
        reference_signs_section = generate_reference_signs_section(
            self.invention_data.reference_signs
        )
        if reference_signs_section:
            sections.append(reference_signs_section)

        # 청구범위
        claims_section = generate_claims_section(self.invention_data.claims)
        if claims_section:
            sections.append(claims_section)

        # 요약서
        sections.append(
            generate_abstract_section(self.invention_data.abstract, self.invention_data)
        )

        self._specification_text = "\n".join(sections)
        return self._specification_text

    def _generate_header(self) -> str:
        """명세서 헤더 생성"""
        lines = []
        lines.append("=" * 60)
        lines.append("특허명세서")
        lines.append("=" * 60)
        lines.append("")
        lines.append(f"발명의 명칭: {self.invention_data.title}")
        lines.append(f"작성일: {datetime.now().strftime('%Y년 %m월 %d일')}")
        lines.append("")

        if self.invention_data.applicants:
            lines.append("【출원인】")
            for applicant in self.invention_data.applicants:
                lines.append(f"  - {applicant.name} ({applicant.nationality})")
            lines.append("")

        if self.invention_data.inventors:
            lines.append("【발명자】")
            for inventor in self.invention_data.inventors:
                lines.append(f"  - {inventor.name} ({inventor.nationality})")
            lines.append("")

        if self.invention_data.priority_claims:
            lines.append("【우선권주장】")
            for priority in self.invention_data.priority_claims:
                lines.append(f"  - {priority}")
            lines.append("")

        lines.append("=" * 60)
        lines.append("")

        return "\n".join(lines)

    def save(
        self,
        output_path: str,
        format: Literal["txt", "md", "docx"] = "txt",
    ) -> str:
        """
        명세서를 파일로 저장

        Args:
            output_path: 출력 파일 경로
            format: 출력 형식 (txt, md, docx)

        Returns:
            저장된 파일 경로
        """
        if self._specification_text is None:
            self.generate()

        output_path = Path(output_path)

        if format == "txt":
            return self._save_as_txt(output_path)
        elif format == "md":
            return self._save_as_markdown(output_path)
        elif format == "docx":
            return self._save_as_docx(output_path)
        else:
            raise ValueError(f"지원하지 않는 형식입니다: {format}")

    def _save_as_txt(self, output_path: Path) -> str:
        """텍스트 파일로 저장"""
        output_path = output_path.with_suffix(".txt")
        output_path.write_text(self._specification_text, encoding="utf-8")
        return str(output_path)

    def _save_as_markdown(self, output_path: Path) -> str:
        """마크다운 파일로 저장"""
        output_path = output_path.with_suffix(".md")

        # 마크다운 형식으로 변환
        md_text = self._convert_to_markdown()
        output_path.write_text(md_text, encoding="utf-8")
        return str(output_path)

    def _convert_to_markdown(self) -> str:
        """텍스트를 마크다운 형식으로 변환"""
        lines = []
        lines.append(f"# 특허명세서: {self.invention_data.title}")
        lines.append("")
        lines.append(f"*작성일: {datetime.now().strftime('%Y년 %m월 %d일')}*")
        lines.append("")

        if self.invention_data.applicants:
            lines.append("## 출원인")
            for applicant in self.invention_data.applicants:
                lines.append(f"- {applicant.name} ({applicant.nationality})")
            lines.append("")

        if self.invention_data.inventors:
            lines.append("## 발명자")
            for inventor in self.invention_data.inventors:
                lines.append(f"- {inventor.name} ({inventor.nationality})")
            lines.append("")

        lines.append("---")
        lines.append("")
        lines.append("# 명세서")
        lines.append("")

        # 발명의 명칭
        lines.append(f"## 발명의 명칭")
        lines.append(f"**{self.invention_data.title}**")
        lines.append("")

        # 기술분야
        lines.append("## 기술분야")
        tf = self.invention_data.technical_field
        lines.append(f"본 발명은 **{tf.main_field}**에 관한 것")
        if tf.sub_fields:
            lines.append(f"으로, 보다 상세하게는 {', '.join(tf.sub_fields)}에 관한 것이다.")
        else:
            lines.append("이다.")
        lines.append("")

        # 발명의 배경이 되는 기술
        lines.append("## 발명의 배경이 되는 기술")
        lines.append(self.invention_data.background_art.conventional_technology)
        lines.append("")

        if self.invention_data.background_art.problems_of_prior_art:
            lines.append("### 종래 기술의 문제점")
            for problem in self.invention_data.background_art.problems_of_prior_art:
                lines.append(f"- {problem}")
            lines.append("")

        # 해결하려는 과제
        lines.append("## 해결하려는 과제")
        lines.append(self.invention_data.problem_to_solve.technical_problem)
        lines.append("")

        if self.invention_data.problem_to_solve.objectives:
            lines.append("### 발명의 목적")
            for obj in self.invention_data.problem_to_solve.objectives:
                lines.append(f"1. {obj}")
            lines.append("")

        # 과제의 해결 수단
        lines.append("## 과제의 해결 수단")
        lines.append(self.invention_data.solution.main_solution)
        lines.append("")

        if self.invention_data.solution.technical_features:
            lines.append("### 기술적 특징")
            for feature in self.invention_data.solution.technical_features:
                lines.append(f"- {feature}")
            lines.append("")

        # 발명의 효과
        lines.append("## 발명의 효과")
        lines.append(self.invention_data.effect.main_effect)
        lines.append("")

        if self.invention_data.effect.additional_effects:
            lines.append("### 부가적 효과")
            for effect in self.invention_data.effect.additional_effects:
                lines.append(f"- {effect}")
            lines.append("")

        # 도면 설명
        if self.invention_data.drawings:
            lines.append("## 도면의 간단한 설명")
            for drawing in self.invention_data.drawings:
                lines.append(
                    f"- **도 {drawing.figure_number}**: {drawing.title} - {drawing.description}"
                )
            lines.append("")

        # 실시예
        if self.invention_data.embodiments:
            lines.append("## 발명을 실시하기 위한 구체적인 내용")
            for emb in self.invention_data.embodiments:
                lines.append(f"### 실시예 {emb.embodiment_number}: {emb.title}")
                lines.append(emb.description)
                lines.append("")

        # 부호의 설명
        if self.invention_data.reference_signs:
            lines.append("## 부호의 설명")
            lines.append("| 부호 | 명칭 |")
            lines.append("|------|------|")
            for ref in self.invention_data.reference_signs:
                lines.append(f"| {ref.sign} | {ref.name} |")
            lines.append("")

        # 청구범위
        if self.invention_data.claims:
            lines.append("## 청구범위")
            for claim in self.invention_data.claims:
                lines.append(f"### 청구항 {claim.claim_number}")
                if claim.claim_type == "dependent" and claim.dependent_on:
                    lines.append(f"*제{claim.dependent_on}항에 종속*")
                lines.append(claim.characterizing_portion)
                lines.append("")

        # 요약서
        lines.append("## 요약서")
        if self.invention_data.abstract:
            lines.append(f"**과제**: {self.invention_data.abstract.problem}")
            lines.append(f"**해결수단**: {self.invention_data.abstract.solution}")
            lines.append(f"**효과**: {self.invention_data.abstract.effect}")
        else:
            lines.append(f"**과제**: {self.invention_data.problem_to_solve.technical_problem}")
            lines.append(f"**해결수단**: {self.invention_data.solution.main_solution}")
            lines.append(f"**효과**: {self.invention_data.effect.main_effect}")
        lines.append("")

        return "\n".join(lines)

    def _save_as_docx(self, output_path: Path) -> str:
        """DOCX 파일로 저장"""
        try:
            from docx import Document
            from docx.shared import Pt, Inches
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            raise ImportError(
                "python-docx 패키지가 필요합니다. 'pip install python-docx'로 설치하세요."
            )

        output_path = output_path.with_suffix(".docx")
        doc = Document()

        # 제목
        title = doc.add_heading("특허명세서", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 발명의 명칭
        doc.add_heading(self.invention_data.title, level=1)
        doc.add_paragraph(f"작성일: {datetime.now().strftime('%Y년 %m월 %d일')}")

        # 출원인/발명자 정보
        if self.invention_data.applicants:
            doc.add_heading("출원인", level=2)
            for applicant in self.invention_data.applicants:
                doc.add_paragraph(f"• {applicant.name} ({applicant.nationality})")

        if self.invention_data.inventors:
            doc.add_heading("발명자", level=2)
            for inventor in self.invention_data.inventors:
                doc.add_paragraph(f"• {inventor.name} ({inventor.nationality})")

        # 명세서 본문
        doc.add_page_break()
        doc.add_heading("명세서", level=1)

        # 기술분야
        doc.add_heading("기술분야", level=2)
        tf = self.invention_data.technical_field
        text = f"본 발명은 {tf.main_field}에 관한 것"
        if tf.sub_fields:
            text += f"으로, 보다 상세하게는 {', '.join(tf.sub_fields)}에 관한 것이다."
        else:
            text += "이다."
        doc.add_paragraph(text)

        # 발명의 배경이 되는 기술
        doc.add_heading("발명의 배경이 되는 기술", level=2)
        doc.add_paragraph(self.invention_data.background_art.conventional_technology)

        if self.invention_data.background_art.problems_of_prior_art:
            doc.add_paragraph("종래 기술의 문제점:")
            for problem in self.invention_data.background_art.problems_of_prior_art:
                doc.add_paragraph(f"• {problem}")

        # 해결하려는 과제
        doc.add_heading("해결하려는 과제", level=2)
        doc.add_paragraph(self.invention_data.problem_to_solve.technical_problem)

        # 과제의 해결 수단
        doc.add_heading("과제의 해결 수단", level=2)
        doc.add_paragraph(self.invention_data.solution.main_solution)

        if self.invention_data.solution.technical_features:
            doc.add_paragraph("기술적 특징:")
            for feature in self.invention_data.solution.technical_features:
                doc.add_paragraph(f"• {feature}")

        # 발명의 효과
        doc.add_heading("발명의 효과", level=2)
        doc.add_paragraph(self.invention_data.effect.main_effect)

        if self.invention_data.effect.additional_effects:
            doc.add_paragraph("부가적 효과:")
            for effect in self.invention_data.effect.additional_effects:
                doc.add_paragraph(f"• {effect}")

        # 청구범위
        if self.invention_data.claims:
            doc.add_page_break()
            doc.add_heading("청구범위", level=1)
            for claim in self.invention_data.claims:
                doc.add_heading(f"청구항 {claim.claim_number}", level=2)
                if claim.claim_type == "dependent" and claim.dependent_on:
                    doc.add_paragraph(f"제{claim.dependent_on}항에 있어서,")
                doc.add_paragraph(claim.characterizing_portion)

        # 요약서
        doc.add_page_break()
        doc.add_heading("요약서", level=1)
        if self.invention_data.abstract:
            doc.add_paragraph(f"[과제] {self.invention_data.abstract.problem}")
            doc.add_paragraph(f"[해결수단] {self.invention_data.abstract.solution}")
            doc.add_paragraph(f"[효과] {self.invention_data.abstract.effect}")
        else:
            doc.add_paragraph(f"[과제] {self.invention_data.problem_to_solve.technical_problem}")
            doc.add_paragraph(f"[해결수단] {self.invention_data.solution.main_solution}")
            doc.add_paragraph(f"[효과] {self.invention_data.effect.main_effect}")

        doc.save(str(output_path))
        return str(output_path)

    def get_statistics(self) -> dict:
        """명세서 통계 정보 반환"""
        if self._specification_text is None:
            self.generate()

        return {
            "title": self.invention_data.title,
            "total_characters": len(self._specification_text),
            "total_lines": self._specification_text.count("\n") + 1,
            "num_claims": len(self.invention_data.claims),
            "num_drawings": len(self.invention_data.drawings),
            "num_embodiments": len(self.invention_data.embodiments),
            "num_reference_signs": len(self.invention_data.reference_signs),
            "num_technical_features": len(self.invention_data.solution.technical_features),
            "num_effects": 1 + len(self.invention_data.effect.additional_effects),
        }
